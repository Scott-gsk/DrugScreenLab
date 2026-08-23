from __future__ import annotations

import hashlib
import json
import sys
import zipfile
from pathlib import Path

import openpyxl
from lxml import etree


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def main(root: Path) -> None:
    files = []
    for path in sorted(root.iterdir()):
        if not path.is_file() or path.suffix in {".html", ".txt"}:
            continue
        files.append({"name": path.name, "bytes": path.stat().st_size, "sha256": sha256(path)})

    workbook_path = root / "NIHMS2149379-supplement-Supplementary_Tables.xlsx"
    with zipfile.ZipFile(workbook_path) as archive:
        bad_member = archive.testzip()
    workbook = openpyxl.load_workbook(workbook_path, read_only=True, data_only=True)
    sheets = []
    for sheet in workbook.worksheets:
        preview = []
        nonempty = 0
        for row_index, row in enumerate(sheet.iter_rows(values_only=True), start=1):
            values = list(row)
            nonempty += sum(value is not None for value in values)
            if row_index <= 12:
                preview.append(values[:20])
        sheets.append(
            {
                "title": sheet.title,
                "max_row": sheet.max_row,
                "max_column": sheet.max_column,
                "nonempty_cells": nonempty,
                "preview_first_12_rows_20_columns": preview,
            }
        )
    workbook.close()

    docx_path = root / "NIHMS2149379-supplement-Supplementary_Materials.docx"
    with zipfile.ZipFile(docx_path) as archive:
        docx_bad_member = archive.testzip()
        document = etree.fromstring(archive.read("word/document.xml"))
    namespaces = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    paragraphs = [
        "".join(paragraph.xpath(".//w:t/text()", namespaces=namespaces)).strip()
        for paragraph in document.xpath(".//w:p", namespaces=namespaces)
    ]
    keywords = (
        "accession", "availability", "repository", "request", "rna", "drug screen", "drug testing",
        "organoid", "cmax", "viability", "dmso", "atp", "celltiter",
    )
    relevant_docx_paragraphs = [
        paragraph for paragraph in paragraphs if paragraph and any(key in paragraph.lower() for key in keywords)
    ]

    workbook = openpyxl.load_workbook(workbook_path, read_only=True, data_only=True)

    def rectangular(title: str) -> list[list[object]]:
        return [list(row) for row in workbook[title].iter_rows(values_only=True)]

    s3 = rectangular("Table S3")
    s5 = rectangular("Table S5")
    s10 = rectangular("Table S10")
    s11 = rectangular("Table S11")
    targeted_samples = [str(value) for value in s3[1][1:] if value is not None]
    whole_transcript_samples = [str(value) for value in s5[1][1:] if value is not None]
    response_samples = [str(row[0]) for row in s11[2:] if row[0] is not None]
    response_drugs = [str(value) for value in s11[1][1:] if value is not None]
    response_observations = sum(
        value is not None for row in s11[2:] if row[0] is not None for value in row[1 : 1 + len(response_drugs)]
    )
    cmax_indexes = [index for index, name in enumerate(response_drugs, start=1) if not name.endswith("DSS")]
    dss_indexes = [index for index, name in enumerate(response_drugs, start=1) if name.endswith("DSS")]
    cmax_observations = sum(
        row[index] is not None for row in s11[2:] if row[0] is not None for index in cmax_indexes
    )
    dss_observations = sum(
        row[index] is not None for row in s11[2:] if row[0] is not None for index in dss_indexes
    )
    dose_curve_rows = [row for row in s10[2:] if row[0] is not None and row[1] is not None]
    pairing = {
        "targeted_expression": {"genes": len(s3) - 2, "samples": len(targeted_samples)},
        "whole_transcript_expression": {"genes": len(s5) - 2, "samples": len(whole_transcript_samples)},
        "pdo_response": {
            "samples": len(response_samples),
            "drug_or_combination_columns": len(response_drugs),
            "nonmissing_observations": response_observations,
            "cmax_columns": len(cmax_indexes),
            "cmax_nonmissing_observations": cmax_observations,
            "dss_derived_columns": len(dss_indexes),
            "dss_nonmissing_observations": dss_observations,
        },
        "dose_curve": {
            "rows": len(dose_curve_rows),
            "unique_samples": len({str(row[0]) for row in dose_curve_rows}),
            "unique_drugs": len({str(row[1]) for row in dose_curve_rows}),
        },
        "targeted_response_overlap": sorted(set(targeted_samples) & set(response_samples)),
        "whole_transcript_response_exact_overlap": sorted(set(whole_transcript_samples) & set(response_samples)),
        "whole_transcript_sample_ids": whole_transcript_samples,
        "response_drugs": response_drugs,
    }
    workbook.close()

    result = {
        "files": files,
        "archive_gate": {"xlsx_bad_member": bad_member, "docx_bad_member": docx_bad_member},
        "supplementary_materials_relevant_paragraphs": relevant_docx_paragraphs,
        "workbook_sheets": sheets,
        "pairing_audit": pairing,
    }
    json.dump(result, sys.stdout, ensure_ascii=False, indent=2, default=str)


if __name__ == "__main__":
    main(Path(sys.argv[1]))
