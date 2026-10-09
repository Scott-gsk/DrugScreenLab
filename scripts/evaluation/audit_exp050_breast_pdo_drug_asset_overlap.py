#!/usr/bin/env python3
"""Coverage-only audit of a public PDO IC50 workbook against measured RNA-response drugs.

This script intentionally reads only worksheet dimensions, PDO identifier labels in
D2:CC2, and drug labels in A4:A52. It never accesses the numeric IC50 response cells.
Run only in the registered WSL2 Conda environment.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import platform
import shutil
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook
import openpyxl


RUN_ID = "20261001_breast_pdo_drug_asset_overlap_v1"
EXPECTED_SHA256 = "b66e558a82f11e607d13019bd3ad7701b081d3a371f20a87989f90031d37aa2b"
EXPECTED_SHAPE = (55, 81)
ROOT = Path(__file__).resolve().parents[2]
INPUT_REL = Path("data/raw/breast_pdo_chen_2021/v1/ADVS-8-2101176-s003.xlsx")
CONFIG_REL = Path("configs/experiments/EXP-050.yaml")
OUTPUT_REL = Path("artifacts/experiments/EXP-050") / RUN_ID

# Frozen name-only coverage aliases. These do not establish chemical identity in
# the Chen workbook, which contains no structure key in the inspected table.
EXPRESSION_DRUGS = [
    {"source_label": "DOX", "coverage_label": "Doxorubicin", "dataset": "GSE287429"},
    {"source_label": "BTZ", "coverage_label": "Bortezomib", "dataset": "GSE287429"},
    {"source_label": "Fulvestrant", "coverage_label": "Fulvestrant", "dataset": "GSE216540"},
    {"source_label": "JQ1", "coverage_label": "JQ1", "dataset": "GSE288798"},
    {"source_label": "MS177", "coverage_label": "MS177", "dataset": "GSE288798"},
]


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def atomic_json(path: Path, value: dict) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_suffix(path.suffix + ".tmp")
    temporary.write_text(json.dumps(value, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    temporary.replace(path)


def normalize_label(value: object) -> str:
    return str(value).strip().casefold()


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--root", type=Path, default=ROOT)
    args = parser.parse_args()
    root = args.root.resolve()
    source = root / INPUT_REL
    config = root / CONFIG_REL
    output_dir = root / OUTPUT_REL

    actual_sha = sha256(source)
    if actual_sha != EXPECTED_SHA256:
        raise SystemExit(f"Registered source SHA-256 mismatch: {actual_sha}")

    workbook = load_workbook(source, read_only=True, data_only=True)
    try:
        if workbook.sheetnames != ["Sheet1"]:
            raise SystemExit(f"Unexpected worksheet names: {workbook.sheetnames}")
        sheet = workbook["Sheet1"]
        shape = (sheet.max_row, sheet.max_column)
        if shape != EXPECTED_SHAPE:
            raise SystemExit(f"Unexpected Sheet1 dimensions: {shape}")

        # The only workbook value reads in this program are these label cells.
        pdo_labels = [
            str(sheet.cell(row=2, column=column).value).strip()
            for column in range(4, 82)
            if sheet.cell(row=2, column=column).value is not None
            and str(sheet.cell(row=2, column=column).value).strip()
        ]
        drug_label_cells = []
        for row in range(4, 53):
            value = sheet.cell(row=row, column=1).value
            if isinstance(value, str) and value.strip():
                drug_label_cells.append({"row": row, "label": value.strip()})
        drug_labels = [entry["label"] for entry in drug_label_cells]
    finally:
        workbook.close()

    if len(pdo_labels) != 78 or len(drug_labels) != 47:
        raise SystemExit(f"Label coverage mismatch: {len(pdo_labels)=}, {len(drug_labels)=}")
    if len({normalize_label(label) for label in drug_labels}) != len(drug_labels):
        raise SystemExit("Duplicate drug labels under case/edge-space normalization")

    chen_by_normalized = {normalize_label(label): label for label in drug_labels}
    overlaps = []
    unmatched = []
    for compound in EXPRESSION_DRUGS:
        matched_label = chen_by_normalized.get(normalize_label(compound["coverage_label"]))
        entry = {**compound, "chen_label": matched_label}
        if matched_label is None:
            unmatched.append(entry)
        else:
            overlaps.append(entry)

    now = datetime.now(timezone.utc).astimezone().isoformat(timespec="seconds")
    summary = {
        "run_id": RUN_ID,
        "status": "COVERAGE_ONLY_COMPLETE",
        "completed_at": now,
        "source": {
            "dataset_id": "chen_2021_breast_pdo_pmc_supplements_v1",
            "file": INPUT_REL.as_posix(),
            "title": "IC50 values (μM) of each drug in the breast cancer organoid lines",
            "bytes": source.stat().st_size,
            "sha256": actual_sha,
            "worksheet": "Sheet1",
            "worksheet_dimensions_rows_columns": list(shape),
            "pdo_identifier_count": len(pdo_labels),
            "unique_pdo_identifier_count": len({normalize_label(x) for x in pdo_labels}),
            "drug_label_count": len(drug_labels),
            "unique_drug_label_count": len({normalize_label(x) for x in drug_labels}),
            "drug_labels": drug_labels,
            "drug_label_cells": drug_label_cells,
            "blank_cells_in_A4_A52": 49 - len(drug_labels),
            "blank_row_indices_in_A4_A52": sorted(set(range(4, 53)) - {entry["row"] for entry in drug_label_cells}),
        },
        "expression_drug_coverage": {
            "measured_transcription_response_drug_count": len(EXPRESSION_DRUGS),
            "coverage_rule": "casefold and trim edge spaces, then exact label match after the frozen source-label to coverage-label mapping",
            "name_level_overlap_count": len(overlaps),
            "name_level_overlaps": overlaps,
            "no_name_level_overlap": unmatched,
            "structure_identity_in_chen_table": "not established; no structure key was read or used",
        },
        "sample_level_crosswalk": {
            "status": "not established",
            "rule": "do not join studies by matching numeric PDO suffixes",
            "interpretation": "78 identifiers belong to the Chen cohort; they do not establish pairing with any GSE expression-response object",
        },
        "access_boundary": {
            "workbook_value_cells_read": ["D2:CC2 PDO labels", "A4:A52 drug labels"],
            "numeric_ic50_response_cells_accessed": False,
            "numeric_ic50_values_read": 0,
            "ic50_statistics_or_ranking": False,
            "sensitivity_bins_or_efficacy_inference": False,
            "gpu_hours": 0.0,
            "target_pdo_inference": False,
        },
    }
    summary_path = output_dir / "summary.json"
    atomic_json(summary_path, summary)
    protocol_snapshot = output_dir / "protocol_snapshot.yaml"
    shutil.copyfile(config, protocol_snapshot)

    manifest = {
        "run_id": RUN_ID,
        "status": summary["status"],
        "completed_at": now,
        "command": "conda run --no-capture-output -n drugscreening-gpu python scripts/evaluation/audit_exp050_breast_pdo_drug_asset_overlap.py",
        "protocol": CONFIG_REL.as_posix(),
        "protocol_sha256": sha256(config),
        "input": {"path": INPUT_REL.as_posix(), "bytes": source.stat().st_size, "sha256": actual_sha},
        "script": {
            "path": Path(__file__).resolve().relative_to(root).as_posix(),
            "sha256": sha256(Path(__file__).resolve()),
        },
        "environment": {
            "platform": platform.platform(),
            "python": platform.python_version(),
            "openpyxl": openpyxl.__version__,
            "gpu_hours": 0.0,
        },
        "outputs": {
            "summary.json": sha256(summary_path),
            "protocol_snapshot.yaml": sha256(protocol_snapshot),
        },
        "preflight_failures": [
            {
                "attempt": 1,
                "cause": "the first script revision omitted the terminal 'b' from the registered workbook SHA-256 constant",
                "stop_point": "checksum check before opening workbook",
                "numeric_ic50_values_read": 0,
                "gpu_hours": 0.0,
            },
            {
                "attempt": 2,
                "cause": "the first label-count assertion expected 49 non-empty drug labels; bounded read found 47",
                "stop_point": "after reading only D2:CC2 and A4:A52 labels, before writing summary outputs",
                "numeric_ic50_values_read": 0,
                "gpu_hours": 0.0,
            },
        ],
    }
    atomic_json(output_dir / "run_manifest.json", manifest)
    print(json.dumps({
        "status": summary["status"],
        "pdo_ids": summary["source"]["pdo_identifier_count"],
        "chen_drug_labels": summary["source"]["drug_label_count"],
        "expression_drugs": len(EXPRESSION_DRUGS),
        "name_level_overlap": [entry["coverage_label"] for entry in overlaps],
        "numeric_ic50_values_read": 0,
        "output": OUTPUT_REL.as_posix(),
    }, ensure_ascii=False))


if __name__ == "__main__":
    main()
