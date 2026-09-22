from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
p=Presentation();p.slide_width=Inches(13.33);p.slide_height=Inches(7.5)
NAVY=RGBColor(23,50,77); BLUE=RGBColor(47,107,138); GRAY=RGBColor(83,97,107); LIGHT=RGBColor(238,244,247)
def slide(title,sub=''):
 s=p.slides.add_slide(p.slide_layouts[6]); sh=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,0,0,p.slide_width,Inches(.16));sh.fill.solid();sh.fill.fore_color.rgb=BLUE;sh.line.fill.background();
 tb=s.shapes.add_textbox(Inches(.55),Inches(.35),Inches(12.2),Inches(.5));tf=tb.text_frame;tf.text=title;tf.paragraphs[0].font.size=Pt(25);tf.paragraphs[0].font.bold=True;tf.paragraphs[0].font.color.rgb=NAVY
 if sub:
  t=s.shapes.add_textbox(Inches(.58),Inches(.88),Inches(12),Inches(.25));t.text=sub;t.text_frame.paragraphs[0].font.size=Pt(11);t.text_frame.paragraphs[0].font.color.rgb=GRAY
 return s
def box(s,x,y,w,h,txt,size=18,bold=False,color=NAVY,fill=None):
 sh=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));
 if fill:
  sh.fill.solid();sh.fill.fore_color.rgb=fill;sh.line.color.rgb=fill
 tf=sh.text_frame;tf.word_wrap=True;tf.margin_left=Inches(.08);tf.margin_right=Inches(.08);tf.text=txt
 for p0 in tf.paragraphs:p0.font.size=Pt(size);p0.font.bold=bold;p0.font.color.rgb=color
 return sh
s=slide('本周结论','类器官单药筛选跨模型迁移')
box(s,.8,1.7,11.8,1.2,'完成第一轮验证：细胞系模型能识别部分表达信号，但当前类器官数据上尚未超过简单基线。',27,True,NAVY,LIGHT)
box(s,.9,3.5,11.5,1.2,'已完成：方案设计｜来源模型验证｜类器官适配测试｜功能数据初测\n下一步：外部验证 + 扩大类器官和药物覆盖',21,False,GRAY)
s=slide('关键进展 1｜完成可验证的筛药方案')
box(s,.7,1.3,12,1.0,'目标：用细胞系大样本药敏信息，辅助小样本类器官对候选单药排序。',22,True,NAVY,LIGHT)
box(s,.8,2.8,3.7,2.3,'输入\n类器官基础表达\n候选药物结构',20,True,NAVY,LIGHT)
box(s,4.8,2.8,3.7,2.3,'输出\n药物敏感性排序\nTop候选药物\n解释依据',20,True,NAVY,LIGHT)
box(s,8.8,2.8,3.7,2.3,'评价规则\n患者留出\n与简单基线比较\n不等同临床疗效',20,True,NAVY,LIGHT)
s=slide('关键进展 2｜来源模型识别到有限表达信号')
box(s,.7,1.25,12,0.6,'动作：按细胞系整组留出，比较“药物平均效应”与“表达 + 药物结构”。',19,False,GRAY)
box(s,.8,2.1,3.7,1.7,'药物平均基线\nNDCG@5 ≈ 0.956',23,True,NAVY,LIGHT)
box(s,4.8,2.1,3.7,1.7,'表达+结构模型\nNDCG@5 ≈ 0.963\n提升约0.007',23,True,NAVY,LIGHT)
box(s,8.8,2.1,3.7,1.7,'表达置换后\n性能下降\n支持使用细胞信号',21,True,NAVY,LIGHT)
box(s,.8,4.5,11.7,.8,'结果含义：来源域存在表达相关信号，但尚未证明可以迁移到类器官。',22,True,NAVY)
s=slide('关键进展 3｜类器官第一轮验证未超过简单基线')
box(s,.7,1.2,12,.6,'动作：在20个类器官、7种严格共享药物上做患者留出比较。',19,False,GRAY)
for i,(a,b) in enumerate([('简单药物基线','≈0.748'),('类器官表达模型','≈0.68–0.73'),('直接迁移模型','≈0.61'),('来源模型适配','≈0.65')]):
 box(s,.75+(i%2)*6.1,2+(i//2)*1.45,5.6,1.05,a+'\n'+b,21,True,NAVY,LIGHT)
box(s,.8,5.25,11.6,.8,'判断：当前主要瓶颈是样本量、共同药物数和跨体系表达差异，暂不继续堆复杂网络。',20,True,NAVY)
s=slide('下周计划与所需支持')
box(s,.7,1.2,12,.55,'计划聚焦“验证能否复现”与“补足数据条件”。',20,True,NAVY,LIGHT)
box(s,.8,2.0,3.8,2.2,'1｜完成外部验证\nPRISM严格未见细胞\n确认表达信号能否跨研究复现',18,True,NAVY,LIGHT)
box(s,4.8,2.0,3.8,2.2,'2｜补充PDO队列\n扩大患者数\n扩大共同药物覆盖',18,True,NAVY,LIGHT)
box(s,8.8,2.0,3.8,2.2,'3｜形成可用基线\n稳定、可解释的类器官药物排序\n再决定是否升级模型',18,True,NAVY,LIGHT)
box(s,.8,5.0,11.7,.8,'需要支持：确认第二个合格PDO队列及表达—单药药敏配对数据来源。',21,True,NAVY)
p.save('本周工作汇报_类器官单药筛选迁移_紧凑版.pptx')
