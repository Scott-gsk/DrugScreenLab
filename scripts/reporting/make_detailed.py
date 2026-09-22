from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
p=Presentation();p.slide_width=Inches(13.33);p.slide_height=Inches(7.5)
NAVY=RGBColor(23,50,77);BLUE=RGBColor(47,107,138);GRAY=RGBColor(83,97,107);LIGHT=RGBColor(238,244,247);GRID=RGBColor(205,218,225)
def base(title,sub=''):
 s=p.slides.add_slide(p.slide_layouts[6]); sh=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,0,0,p.slide_width,Inches(.16));sh.fill.solid();sh.fill.fore_color.rgb=BLUE;sh.line.fill.background()
 t=s.shapes.add_textbox(Inches(.55),Inches(.35),Inches(12.2),Inches(.5));t.text=title;t.text_frame.paragraphs[0].font.size=Pt(25);t.text_frame.paragraphs[0].font.bold=True;t.text_frame.paragraphs[0].font.color.rgb=NAVY
 if sub:
  q=s.shapes.add_textbox(Inches(.58),Inches(.88),Inches(12),Inches(.25));q.text=sub;q.text_frame.paragraphs[0].font.size=Pt(11);q.text_frame.paragraphs[0].font.color.rgb=GRAY
 return s
def tx(s,x,y,w,h,text,size=18,bold=False,color=NAVY,fill=None):
 sh=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));
 if fill: sh.fill.solid();sh.fill.fore_color.rgb=fill;sh.line.color.rgb=fill
 tf=sh.text_frame;tf.word_wrap=True;tf.margin_left=Inches(.08);tf.margin_right=Inches(.08);tf.text=text
 for p0 in tf.paragraphs:p0.font.size=Pt(size);p0.font.bold=bold;p0.font.color.rgb=color
 return sh
def tbl(s,rows,x,y,w,h,fs=14):
 sh=s.shapes.add_table(len(rows),len(rows[0]),Inches(x),Inches(y),Inches(w),Inches(h));tab=sh.table
 widths=[w/len(rows[0])]*len(rows[0])
 if len(rows[0])==3: widths=[w*.22,w*.38,w*.40]
 for j,ww in enumerate(widths): tab.columns[j].width=Inches(ww)
 for i,row in enumerate(rows):
  for j,val in enumerate(row):
   c=tab.cell(i,j);c.text=str(val);c.margin_left=Inches(.06);c.margin_right=Inches(.06);c.margin_top=Inches(.03);c.margin_bottom=Inches(.03);c.fill.solid();c.fill.fore_color.rgb=BLUE if i==0 else (LIGHT if i%2 else RGBColor(255,255,255))
   for p0 in c.text_frame.paragraphs:p0.font.size=Pt(fs);p0.font.color.rgb=RGBColor(255,255,255) if i==0 else NAVY;p0.font.bold=(i==0);p0.alignment=PP_ALIGN.LEFT
 return sh
s=base('本周结论','类器官单药筛选跨模型迁移')
tx(s,.8,1.55,11.7,1.25,'完成第一轮验证：细胞系模型能识别部分表达信号，但在当前类器官队列上尚未超过简单基线。',27,True,NAVY,LIGHT)
tbl(s,[['本周已完成','当前判断'],['方案设计、来源模型验证、类器官适配、功能数据初测','模型链路已打通，但迁移有效性尚未成立'],['PRISM身份与药物覆盖审计','具备开展跨研究外部验证的条件']],.8,3.35,11.7,1.7,16)
tx(s,.8,5.55,11.7,.65,'下周重点：外部验证 + 补充PDO数据 + 固化可靠基线。',21,True,NAVY)
s=base('关键进展 1｜完成可验证的筛药方案','动作 + 结果 + 对项目的意义')
tbl(s,[['本周动作','形成的具体内容','带来的结果'],['明确研究问题','用细胞系药敏信息辅助PDO候选单药排序','研究目标从“训练模型”转为“支持药物优先级判断”'],['固定输入输出','输入：基线表达、药物结构；输出：药物排序、Top候选、解释依据','后续不同模型可在同一任务上公平比较'],['固定评价规则','患者留出；以患者内排序为主；与简单基线比较','降低数据泄漏和药物普遍效应造成的虚高'],['明确证据边界','研究级预筛选，不直接等同临床疗效','避免将算法结果外推为患者治疗结论']],.65,1.35,12.05,4.35,13)
s=base('关键进展 2｜来源模型识别到有限表达信号','细胞系内部验证结果')
tbl(s,[['比较项','结果','说明'],['药物平均敏感性基线','NDCG@5 ≈ 0.956','只利用药物的普遍作用强弱'],['表达 + 药物结构模型','NDCG@5 ≈ 0.963','相对提升约0.007'],['细胞内相关性','Spearman提升约0.03','对同一细胞的药物排序更贴近实测结果'],['表达反事实置换','同癌种内置换后下降；泛癌置换下降更明显','模型使用了部分细胞表达和癌种信息'],['本周结论','来源域存在有限表达信号','尚不能据此证明可以迁移到PDO']],.65,1.35,12.05,4.65,13)
s=base('关键进展 3｜第一轮PDO验证未超过简单基线','20个PDO｜严格匹配7种共同药物｜患者留出评估')
tbl(s,[['模型/方法','排序结果','结果含义'],['简单药物平均敏感性','约0.748','当前最强基线'],['PDO自身表达模型','约0.68–0.73','小样本下未提取稳定个体化信号'],['直接使用细胞系模型','约0.61','出现明显负迁移'],['来源模型适配','约0.65','低自由度适配暂未解决负迁移'],['数据限制','20个PDO、7种共同药物','只能作为开发队列，不能作为独立外部验证']],.65,1.3,12.05,4.75,13)
tx(s,.8,6.25,11.7,.45,'判断：当前瓶颈优先指向样本量、共同药物覆盖和表达尺度差异。',18,True,NAVY)
s=base('关键进展 4｜功能数据暂未改善最终筛药结果','三维表达和类器官功能依赖的辅助训练')
tbl(s,[['本周动作','观察到的结果','当前处理'],['加入三维表达辅助任务','相对基础训练有变化，但未超过简单基线','保留为辅助信息，不作为主模型依据'],['加入类器官功能依赖监督','相对关闭辅助任务略有改善','尚不足以证明改善药物排序'],['功能标签置换对照','真实标签略优于置换结果','提示可能有方向性信号，但证据不足'],['PRISM外部验证准备','完成877→354个严格未见细胞筛选；104个共同药物','下一步执行正式跨研究测试']],.65,1.3,12.05,4.7,13)
s=base('下周计划与所需支持','每项计划对应一个可交付结果')
tbl(s,[['计划','下周交付','需要支持'],['完成PRISM跨研究验证','确认表达信号能否在另一药敏研究的未见细胞上复现','维持当前计算资源和数据版本'],['补充第二个PDO队列','增加患者数，并扩大严格共同药物数量','协助确认可公开获取的表达—单药药敏配对数据'],['固化可靠PDO基线','形成稳定、可解释的类器官候选药物排序流程','确认后续优先癌种和候选药物范围'],['形成继续/停止判断','若外部仍无增益，转向基线和负迁移边界；若有增益，再扩展通路和功能信息','确认下一阶段资源投入优先级']],.65,1.25,12.05,4.8,13)
tx(s,.8,6.3,11.7,.45,'需要领导决策：是否支持获取第二个具备“表达 + 单药响应”配对数据的PDO队列。',18,True,NAVY,LIGHT)
p.save('本周工作汇报_类器官单药筛选迁移_表格详版.pptx')
