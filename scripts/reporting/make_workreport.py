from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
p=Presentation();p.slide_width=Inches(13.33);p.slide_height=Inches(7.5)
NAVY=RGBColor(23,50,77);BLUE=RGBColor(47,107,138);GRAY=RGBColor(83,97,107);LIGHT=RGBColor(238,244,247)
def base(title):
 s=p.slides.add_slide(p.slide_layouts[6]);sh=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,0,0,p.slide_width,Inches(.16));sh.fill.solid();sh.fill.fore_color.rgb=BLUE;sh.line.fill.background();t=s.shapes.add_textbox(Inches(.55),Inches(.35),Inches(12.2),Inches(.5));t.text=title;t.text_frame.paragraphs[0].font.size=Pt(25);t.text_frame.paragraphs[0].font.bold=True;t.text_frame.paragraphs[0].font.color.rgb=NAVY;return s
def tx(s,x,y,w,h,text,size=17,bold=False,color=NAVY,fill=None):
 sh=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));
 if fill:sh.fill.solid();sh.fill.fore_color.rgb=fill;sh.line.color.rgb=fill
 tf=sh.text_frame;tf.word_wrap=True;tf.margin_left=Inches(.08);tf.margin_right=Inches(.08);tf.text=text
 for q in tf.paragraphs:q.font.size=Pt(size);q.font.bold=bold;q.font.color.rgb=color
 return sh
def tbl(s,rows,x,y,w,h,fs=13):
 sh=s.shapes.add_table(len(rows),len(rows[0]),Inches(x),Inches(y),Inches(w),Inches(h));tab=sh.table
 widths=[w*.22,w*.38,w*.40] if len(rows[0])==3 else [w/len(rows[0])]*len(rows[0])
 for j,ww in enumerate(widths):tab.columns[j].width=Inches(ww)
 for i,row in enumerate(rows):
  for j,val in enumerate(row):
   c=tab.cell(i,j);c.text=str(val);c.margin_left=Inches(.06);c.margin_right=Inches(.06);c.fill.solid();c.fill.fore_color.rgb=BLUE if i==0 else (LIGHT if i%2 else RGBColor(255,255,255))
   for q in c.text_frame.paragraphs:q.font.size=Pt(fs);q.font.color.rgb=RGBColor(255,255,255) if i==0 else NAVY;q.font.bold=(i==0)
#1
s=base('1｜本周工作概况：完成跨模型迁移的第一轮验证')
tx(s,.7,1.05,12,.85,'本周围绕“细胞系药敏模型能否帮助类器官单药筛选”开展方案设计、来源域验证、PDO适配实验和外部验证准备，形成了从数据准入到结果判定的完整实验链路。',19,False,GRAY)
tbl(s,[['工作模块','本周完成的工作','阶段产出'],['研究方案','明确研究背景、输入输出、数据角色和判定标准','形成可执行的细胞系→PDO迁移方案'],['来源模型','完成细胞系整组留出和表达置换验证','确认来源域存在有限表达相关信号'],['类器官迁移','完成20个PDO、7种共同药物的患者留出比较','发现当前迁移未超过简单基线'],['外部验证准备','完成PRISM身份排除和药物覆盖审计','获得354个严格未见细胞、104种共同药物']],.7,2.15,11.95,3.2,13)
tx(s,.7,5.75,12,.6,'本周结论：模型链路已打通，但类器官迁移有效性尚未成立，当前瓶颈集中在数据规模和跨体系差异。',20,True,NAVY,LIGHT)
#2
s=base('2｜实验背景与研究目的')
tx(s,.7,1.05,12,.9,'细胞系药敏数据规模大，适合训练模型；类器官更接近患者肿瘤，但样本和药物数据有限。本周方案的核心目的是判断：细胞系中学习到的药物反应规律，能否帮助类器官对候选单药进行排序。',18,False,GRAY)
tbl(s,[['内容','方案设计','为什么需要'],['研究对象','PDO/PDxO单药筛选，不涉及联合用药','保持任务聚焦，避免将单药结果外推到联合治疗'],['模型输入','类器官基线表达 + 候选药物结构','使用患者样本可获得的信息进行预测'],['模型输出','药物敏感性排序、Top候选和解释依据','支持后续实验优先级安排'],['主要评价','患者留出下的类器官内药物排序','避免同一患者信息泄漏，检验真实个体化信号'],['结果边界','研究级预筛选，不等同临床疗效','明确当前算法的使用范围']],.7,2.15,11.95,3.55,13)
#3
s=base('3｜来源模型实验：确认细胞系中存在表达相关信号')
tx(s,.7,1.05,12,.8,'为避免直接跳到类器官结论，本周先在细胞系内部按生物模型整组留出，比较只依赖药物平均效应的方法和加入细胞基础表达的方法，并通过表达置换检查模型是否真正使用了细胞信息。',18,False,GRAY)
tbl(s,[['实验动作','结果','解释'],['细胞系整组留出','药物平均基线NDCG@5≈0.956；表达+药物结构模型≈0.963','表达信息带来约0.007提升'],['多初始化重复','三个初始化方向一致','提升不是单次随机训练造成'],['同癌种内表达置换','性能下降','模型使用了癌种内细胞表达差异'],['泛癌表达置换','下降更明显','癌种构成和表达信息共同贡献'],['阶段结论','来源域存在有限表达信号','尚不能推断该信号能迁移到PDO']],.7,2.15,11.95,3.55,13)
#4
s=base('4｜类器官实验：第一轮迁移未超过简单基线')
tx(s,.7,1.05,12,.85,'随后在Zhu类器官队列开展患者外五折验证。该队列只有20个PDO，严格结构匹配后仅有7种共同药物，因此本轮用于开发和负迁移诊断，不能作为独立外部验证。',18,False,GRAY)
tbl(s,[['比较方法','排序结果','本周判断'],['药物平均敏感性','约0.748','当前最强基线'],['类器官自身表达模型','约0.68–0.73','小样本下未形成稳定个体化信号'],['直接使用细胞系模型','约0.61','出现明显负迁移'],['来源模型适配','约0.65','适配未能抵消体系差异'],['功能依赖辅助训练','相对关闭辅助任务略有改善，但仍低于基线','暂不扩大复杂功能网络']],.7,2.15,11.95,3.55,13)
#5
s=base('5｜本周结论与下周计划')
tx(s,.7,1.05,12,.9,'本周结果把项目从“尝试复杂模型”推进到“明确验证边界”：来源域表达信号已经观察到，但在当前PDO数据条件下尚未形成可迁移的筛药增益。下一步优先补足独立数据和外部验证条件。',18,False,GRAY)
tbl(s,[['下周计划','预期结果','需要支持'],['完成PRISM跨研究测试','判断表达信号能否在另一药敏研究的未见细胞上复现','维持数据版本和计算资源'],['寻找第二个PDO队列','增加患者数和严格共同药物覆盖','协助确认表达+单药响应配对数据来源'],['固化可靠基线模型','形成稳定、可解释的类器官候选药物排序流程','确认优先癌种和候选药物范围'],['阶段决策','外部仍阴性则转向基线和负迁移边界；出现增益再扩展通路/功能模型','确认后续资源投入优先级']],.7,2.15,11.95,3.55,13)
tx(s,.7,6.05,12,.55,'需要领导支持：确认第二个具备“基线表达 + 单药响应”配对数据的PDO队列。',19,True,NAVY,LIGHT)
p.save('本周工作汇报_类器官单药筛选迁移_工作汇报版.pptx')
