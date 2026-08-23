"""Leakage-safe EXP-010 public/local input parsers."""

from __future__ import annotations

import json
import csv
import gzip
import io
import re
import tarfile
import zipfile
import urllib.parse
import urllib.request
from concurrent.futures import ThreadPoolExecutor, as_completed
from pathlib import Path

import numpy as np
import pandas as pd

from drug_screen.data.exp010_public import ordered_gene_universe, sha256


COMBINATION_MARKERS = ("&", "+", "FLIRI", "FLOX", "FLV", "MAPK.3X", "VOLASERTIB.SN38", "TAS-102")
UNSUPPORTED_MODALITIES = {"CETUXIMAB"}  # antibody; PharmaFormer requires a small-molecule SMILES
PUBCHEM_CID_OVERRIDES = {"BMS.754807": "24785538"}


def _patient_id(sample_id: str) -> str:
    match = re.match(r"^(Pt[^_]+)", str(sample_id), flags=re.IGNORECASE)
    if not match:
        raise ValueError(f"cannot derive patient identity from {sample_id!r}")
    return match.group(1)


def parse_mcrc(root: Path) -> tuple[pd.DataFrame, pd.DataFrame, dict[str, object]]:
    """Return log2(CPM+1) expression (samples x genes) and monotherapy DSS."""
    expression_raw = pd.read_excel(root / "Data S2.xlsx")
    required = {"ensembl_gene_id", "hgnc_symbol"}
    if not required.issubset(expression_raw.columns):
        raise ValueError(f"mCRC S2 missing {required - set(expression_raw.columns)}")
    sample_columns = [c for c in expression_raw.columns if c not in {"ensembl_gene_id", "entrezgene_id", "hgnc_symbol"}]
    symbols = expression_raw["hgnc_symbol"].astype("string").str.strip().str.upper()
    counts = expression_raw[sample_columns].apply(pd.to_numeric, errors="coerce").fillna(0.0)
    counts.index = symbols
    counts = counts.loc[counts.index.notna() & (counts.index != "")].groupby(level=0, sort=True).sum()
    library = counts.sum(axis=0)
    if (library <= 0).any():
        raise ValueError("mCRC expression contains an empty library")
    expression = np.log2(counts.divide(library, axis=1) * 1_000_000 + 1).T
    expression.index.name = "sample_id"

    wide = pd.read_excel(root / "Data S5.xlsx")
    drug_columns = [c for c in wide.columns if not str(c).startswith("run_id")]
    drug_columns.remove("sample_id")
    long = wide.melt(id_vars="sample_id", value_vars=drug_columns, var_name="source_drug", value_name="dss").dropna(subset=["dss"])
    long["drug_name"] = long["source_drug"].str.replace(r"_lib[12]$", "", regex=True)
    long["is_combination"] = long["drug_name"].str.upper().map(lambda x: any(m in x for m in COMBINATION_MARKERS))
    long["unsupported_modality"] = long["drug_name"].str.upper().isin(UNSUPPORTED_MODALITIES)
    long = long.loc[~long["is_combination"] & ~long["unsupported_modality"]].copy()
    long["patient_id"] = long["sample_id"].map(_patient_id)
    # Repeated library measurements are technical repeats; median is frozen.
    responses = long.groupby(["patient_id", "sample_id", "drug_name"], as_index=False, sort=True)["dss"].median()
    responses["higher_is_sensitive"] = True
    common_samples = sorted(set(expression.index) & set(responses["sample_id"]))
    expression = expression.loc[common_samples]
    responses = responses.loc[responses["sample_id"].isin(common_samples)].reset_index(drop=True)
    per_sample = responses.groupby("sample_id")["drug_name"].nunique()
    summary = {
        "expression_samples": len(expression), "genes": expression.shape[1],
        "response_rows": len(responses), "patients": responses["patient_id"].nunique(),
        "pdo_samples": responses["sample_id"].nunique(), "monotherapies": responses["drug_name"].nunique(),
        "min_drugs_per_pdo": int(per_sample.min()), "samples_at_least_5_drugs": int((per_sample >= 5).sum()),
        "expression_source_sha256": sha256(root / "Data S2.xlsx"),
        "response_source_sha256": sha256(root / "Data S5.xlsx"),
        "normalization": "log2(CPM+1) from integer-like counts; not TPM",
        "technical_repeat_rule": "median after excluding combinations",
    }
    if summary["patients"] < 40 or summary["min_drugs_per_pdo"] < 5:
        raise ValueError(f"mCRC hard gate failed: {summary}")
    return expression, responses, summary


def parse_gdsc_fitted(path: Path) -> tuple[pd.DataFrame, dict[str, object]]:
    """Read the official malformed-dimension workbook without trusting A1 metadata."""
    from openpyxl import load_workbook
    sheet = load_workbook(path, read_only=True, data_only=True).active
    sheet.reset_dimensions()
    rows = sheet.iter_rows(values_only=True)
    header = next(rows)
    frame = pd.DataFrame(rows, columns=header)
    required = {"COSMIC_ID", "SANGER_MODEL_ID", "DRUG_ID", "DRUG_NAME", "LN_IC50", "AUC"}
    if not required.issubset(frame.columns):
        raise ValueError(f"GDSC fitted missing {required - set(frame.columns)}")
    summary = {
        "rows": len(frame), "cell_lines": int(frame["SANGER_MODEL_ID"].nunique()),
        "drugs": int(frame["DRUG_ID"].nunique()), "sha256": sha256(path),
        "paper_pair_count": 87_596,
        "is_exact_paper_pair_reconstruction": len(frame) == 87_596,
    }
    return frame, summary


def audit_gdsc_raw(path: Path) -> dict[str, object]:
    with path.open("r", encoding="utf-8", newline="") as handle:
        header = next(csv.reader(handle))
    required = {"SANGER_MODEL_ID", "DRUG_ID", "CONC", "INTENSITY"}
    if not required.issubset(header):
        raise ValueError(f"GDSC raw missing columns {required - set(header)}")
    return {"bytes": path.stat().st_size, "sha256": sha256(path), "columns": header, "schema_ok": True}


def parse_gse303509(path: Path) -> tuple[pd.DataFrame, pd.DataFrame, dict[str, object]]:
    columns: dict[str, pd.Series] = {}
    contexts: list[dict[str, str]] = []
    with tarfile.open(path) as archive:
        members = [m for m in archive.getmembers() if m.isfile() and m.name.endswith(".genes.csv.gz")]
        for member in members:
            match = re.match(r"(?P<gsm>GSM\d+)_T(?P<torg>\d+)_", Path(member.name).name)
            if not match:
                raise ValueError(f"unrecognized GSE303509 member: {member.name}")
            handle = archive.extractfile(member)
            if handle is None:
                raise ValueError(f"cannot read {member.name}")
            with gzip.GzipFile(fileobj=handle) as compressed:
                frame = pd.read_csv(compressed, usecols=["gene_id", "TPM"])
            gene = frame["gene_id"].str.replace(r"\.\d+$", "", regex=True)
            values = pd.Series(frame["TPM"].to_numpy(float), index=gene).groupby(level=0).sum()
            sample_id = "TORG" + match.group("torg")
            columns[sample_id] = values
            contexts.append({"sample_id": sample_id, "geo_sample_id": match.group("gsm"), "tissue": "TNBC organoid", "assay": "RSEM gene TPM"})
    matrix = pd.DataFrame(columns).T.sort_index()
    matrix.index.name = "sample_id"
    context = pd.DataFrame(contexts).sort_values("sample_id").reset_index(drop=True)
    summary = {"samples": len(matrix), "genes_union": matrix.shape[1], "complete_values": int(matrix.notna().all(axis=0).sum()), "sha256": sha256(path)}
    if summary["samples"] != 14:
        raise ValueError(f"GSE303509 expected 14 samples, got {summary['samples']}")
    return matrix, context, summary


def parse_gse287429(path: Path) -> tuple[pd.DataFrame, pd.DataFrame, dict[str, object]]:
    matrix = pd.read_csv(path, compression="gzip", index_col="GeneID").T
    pattern = re.compile(r"^(CTR|DOX|BTZ)(23|47)h_([12])$")
    contexts = []
    for sample in matrix.index:
        match = pattern.match(str(sample))
        if not match:
            raise ValueError(f"GSE287429 unexpected sample: {sample}")
        treatment, hours, replicate = match.groups()
        contexts.append({"sample_id": sample, "treatment": treatment, "time_hours": int(hours), "replicate": int(replicate), "control_group": f"CTR{hours}h"})
    context = pd.DataFrame(contexts)
    summary = {"samples": len(matrix), "genes": matrix.shape[1], "contrasts": ["DOX23h_vs_CTR23h", "BTZ23h_vs_CTR23h", "DOX47h_vs_CTR47h", "BTZ47h_vs_CTR47h"], "sha256": sha256(path), "input_claim": "author-normalized counts; not raw counts"}
    if len(matrix) != 12 or context.groupby(["treatment", "time_hours"]).size().ne(2).any():
        raise ValueError(f"GSE287429 matched-control gate failed: {summary}")
    return matrix, context, summary


def parse_cmp_tpm(path: Path, model_ids: set[str]) -> tuple[pd.DataFrame, dict[str, object]]:
    member = "rnaseq_tpm_20220624.csv"
    with zipfile.ZipFile(path) as archive:
        if member not in archive.namelist():
            raise ValueError(f"CMP archive missing {member}")
        with archive.open(member) as handle:
            header = handle.readline().decode("utf-8").rstrip("\r\n").split(",")
        selected_positions = [i for i, value in enumerate(header) if value in model_ids]
        usecols = [0, 1, *selected_positions]
        with archive.open(member) as handle:
            raw = pd.read_csv(handle, header=None, skiprows=5, usecols=usecols, low_memory=False)
    raw.columns = ["gene_id", "symbol", *[header[i] for i in selected_positions]]
    symbols = raw.pop("symbol").astype("string").str.strip().str.upper()
    raw = raw.drop(columns="gene_id").apply(pd.to_numeric, errors="coerce")
    raw.index = symbols
    raw = raw.loc[raw.index.notna() & (raw.index != "")].groupby(level=0, sort=True).sum(min_count=1)
    matrix = raw.T
    matrix.index.name = "model_id"
    summary = {"models_requested": len(model_ids), "models_mapped": len(matrix), "genes": matrix.shape[1], "tpm_member": member, "archive_sha256": sha256(path)}
    if len(matrix) / max(len(model_ids), 1) < 0.9:
        raise ValueError(f"CMP-to-GDSC model identity gate failed: {summary}")
    return matrix, summary


def align_expression_contract(
    mcrc: pd.DataFrame, cmp: pd.DataFrame, gse303509: pd.DataFrame, mcrc_source: Path,
) -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame, dict[str, object]]:
    mapping = pd.read_excel(mcrc_source, usecols=["ensembl_gene_id", "hgnc_symbol"])
    mapping["ensembl_gene_id"] = mapping["ensembl_gene_id"].astype("string").str.replace(r"\.\d+$", "", regex=True)
    mapping["hgnc_symbol"] = mapping["hgnc_symbol"].astype("string").str.strip().str.upper()
    ensembl_to_symbol = mapping.dropna().drop_duplicates("ensembl_gene_id").set_index("ensembl_gene_id")["hgnc_symbol"]
    symbols = pd.Index(gse303509.columns).map(ensembl_to_symbol)
    converted = gse303509.copy()
    converted.columns = symbols
    converted = converted.loc[:, converted.columns.notna()]
    converted = converted.T.groupby(level=0, sort=True).sum().T
    genes, digest = ordered_gene_universe(mcrc.columns, cmp.columns, converted.columns, minimum=15_000)
    summary = {"ordered_genes": len(genes), "gene_universe_sha256": digest, "order": "lexicographic uppercase HGNC symbol"}
    return mcrc.loc[:, genes], cmp.loc[:, genes], converted.loc[:, genes], summary


def fetch_pubchem(names: list[str], cache: Path, max_workers: int = 6) -> tuple[pd.DataFrame, dict[str, object]]:
    """Resolve monotherapy names through PubChem PUG REST, caching every result."""
    cache.parent.mkdir(parents=True, exist_ok=True)
    existing = {r["primary_name"]: r for r in json.loads(cache.read_text())} if cache.exists() else {}
    requested = sorted(set(names), key=str.casefold)

    def resolve(name: str) -> dict[str, object]:
        identifier = PUBCHEM_CID_OVERRIDES.get(name, name)
        namespace = "cid" if name in PUBCHEM_CID_OVERRIDES else "name"
        encoded = urllib.parse.quote(identifier, safe="")
        url = f"https://pubchem.ncbi.nlm.nih.gov/rest/pug/compound/{namespace}/{encoded}/property/CanonicalSMILES,InChIKey/JSON"
        attempts: list[dict[str, object]] = []
        for attempt in range(1, 3):
            try:
                with urllib.request.urlopen(url, timeout=15) as response:
                    payload = json.load(response)
                prop = payload["PropertyTable"]["Properties"][0]
                attempts.append({"attempt": attempt, "status": 200, "identifier": identifier})
                return {"primary_name": name, "source_url": url, "pubchem_cid": str(prop["CID"]), "canonical_smiles": prop.get("ConnectivitySMILES") or prop.get("CanonicalSMILES"), "inchikey": prop.get("InChIKey"), "mapped": True, "attempts": attempts}
            except Exception as error:
                attempts.append({"attempt": attempt, "status": getattr(error, "code", None), "error": type(error).__name__, "identifier": identifier})
        return {"primary_name": name, "source_url": url, "pubchem_cid": None, "canonical_smiles": None, "inchikey": None, "mapped": False, "attempts": attempts}

    records_by_name = {name: existing[name] for name in requested if existing.get(name, {}).get("mapped")}
    pending = [name for name in requested if name not in records_by_name]
    with ThreadPoolExecutor(max_workers=max_workers) as pool:
        futures = {pool.submit(resolve, name): name for name in pending}
        for future in as_completed(futures):
            record = future.result()
            records_by_name[record["primary_name"]] = record
            current = [records_by_name[n] for n in requested if n in records_by_name]
            temporary = cache.with_suffix(cache.suffix + ".tmp")
            temporary.write_text(json.dumps(current, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
            temporary.replace(cache)
    records = [records_by_name[name] for name in requested]
    frame = pd.DataFrame(records)
    mapped = int(frame["mapped"].sum()) if len(frame) else 0
    summary = {"drugs": len(frame), "mapped": mapped, "mapping_rate": mapped / len(frame) if len(frame) else 0.0}
    return frame, summary
